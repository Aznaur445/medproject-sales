"""Builds the default proposal template app/docx_templates/default.docx (docxtpl / Jinja tags).

Run: uv run python scripts/build_proposal_template.py
The result is an ordinary Word file: it can be opened and restyled in Word as long as the {{ }} / {% %} tags stay.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1B, 0x4F, 0x72)
MUTED = RGBColor(0x66, 0x70, 0x85)
OUT = Path(__file__).resolve().parent.parent / "app" / "docx_templates" / "default.docx"


def shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


def para(doc_or_cell, text: str = "", *, bold=False, size=10.5, color=None, align=None, space_after=4, italic=False):
    p = doc_or_cell.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text: str) -> None:
    p = para(doc, text, bold=True, size=13, color=ACCENT, space_after=6)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)

    footer = section.footer.paragraphs[0]
    footer.text = "{{ req.short_name }} · ИНН {{ req.inn }} · {{ req.website }}"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED
    header = section.header.paragraphs[0]
    header.text = "{{ req.brand|upper }}    КП для {{ customer_name }}"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED

    # Title block
    top = doc.add_table(rows=1, cols=2)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = top.rows[0].cells
    left.paragraphs[0].add_run("{{ req.brand|upper }}").bold = True
    left.paragraphs[0].runs[0].font.size = Pt(20)
    left.paragraphs[0].runs[0].font.color.rgb = ACCENT
    para(left, "{{ req.tagline }}", size=9, color=MUTED)
    for label, value in (
        ("Заказчик:", "{{ customer_name }}"),
        ("Исполнитель:", "{{ req.legal_name }}"),
        ("ИНН / ОГРНИП:", "{{ req.inn }} / {{ req.ogrnip }}"),
        ("Сайт:", "{{ req.website }}"),
    ):
        p = right.paragraphs[0] if label == "Заказчик:" else right.add_paragraph()
        r = p.add_run(label + " ")
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        p.add_run(value).font.size = Pt(9)
        p.paragraph_format.space_after = Pt(1)

    para(doc, "", space_after=6)
    para(
        doc,
        "ТЕХНИКО-КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ",
        bold=True,
        size=16,
        color=ACCENT,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    para(
        doc,
        "на выполнение проектных работ · {{ total_text }}",
        size=11,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )

    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in (("Объект", "{{ object_name }}"), ("Адрес", "{{ address }}"), ("Площадь", "{{ area_text }}")):
        row = info.add_row().cells
        row[0].text, row[1].text = label, value
        shade(row[0], "EEF2F6")
        row[0].paragraphs[0].runs[0].bold = True
        row[0].width, row[1].width = Cm(4), Cm(13)

    para(doc, "", space_after=4)
    para(doc, "Уважаемые коллеги!", bold=True, space_after=6)
    para(doc, "{{ intro }}", space_after=6).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para(doc, "{% if req.okved_note %}{{ req.okved_note }}{% endif %}", space_after=6)

    heading(doc, "1. Состав документации и стоимость")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("№", "Наименование работ / раздел", "Стоимость, ₽"), strict=True):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        shade(cell, "1B4F72")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    start = table.add_row().cells
    start[0].text = "{%tr for r in rows %}"
    body = table.add_row().cells
    body[0].text = "{{ r.num }}"
    body[1].text = "{{r r.title }}"
    body[2].text = "{{r r.price }}"
    body[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    end = table.add_row().cells
    end[0].text = "{%tr endfor %}"
    total = table.add_row().cells
    merged = total[0].merge(total[1])
    merged.text = "ОБЩАЯ СТОИМОСТЬ"
    merged.paragraphs[0].runs[0].bold = True
    total[2].text = "{{ total_text }}"
    total[2].paragraphs[0].runs[0].bold = True
    total[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in table.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Cm(1.4), Cm(12.1), Cm(3.5)
    para(doc, "{{ vat_note }}", size=9, italic=True, color=MUTED, space_after=6)

    heading(doc, "2. Сроки и условия оплаты")
    para(doc, "Срок выполнения работ: {{ duration }}.")
    para(doc, "Порядок оплаты{% if payment_options|length > 1 %} на выбор заказчика{% endif %}:")
    para(doc, "{%p for opt in payment_options %}")
    para(doc, "— {{ opt }}")
    para(doc, "{%p endfor %}")
    para(doc, "Документация передаётся поэтапно, по мере готовности разделов, с оформлением актов сдачи-приёмки.")

    heading(doc, "3. Гарантии качества и организация работы")
    para(doc, "{{ quality }}").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para(doc, "{% if req.gip_note %}{{ req.gip_note }}{% endif %}")

    heading(doc, "4. Преимущества работы с «{{ req.brand }}»")
    para(doc, "{%p for a in advantages %}")
    para(doc, "— {{ a }}")
    para(doc, "{%p endfor %}")

    para(doc, "{%p if cases %}")
    heading(doc, "5. Реализованные проекты")
    para(doc, "{%p for c in cases %}")
    para(doc, "— {{ c }}")
    para(doc, "{%p endfor %}")
    para(doc, "{%p endif %}")

    para(doc, "", space_after=6)
    para(doc, "Предложение действительно до {{ valid_until }}.", italic=True, color=MUTED)
    para(doc, "", space_after=6)

    sign = doc.add_table(rows=1, cols=2)
    sign_left, sign_right = sign.rows[0].cells
    sign_left.paragraphs[0].add_run("{{ signature }}")
    para(sign_left, "{{ req.legal_name }}", bold=True)
    sign_right.paragraphs[0].add_run("Контакты").bold = True
    para(sign_right, "Тел.: {{ req.phone }}", size=9.5, space_after=1)
    para(sign_right, "E-mail: {{ req.email }}", size=9.5, space_after=1)
    para(sign_right, "Сайт: {{ req.website }}", size=9.5, space_after=1)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
