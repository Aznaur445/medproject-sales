import io
import shutil
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
from docx import Document

from app.services.proposal_doc import ProposalContent, ProposalSection, docx_to_pdf, money, render_docx
from app.services.settings_store import ProposalDefaults, Requisites

REQ = Requisites(
    legal_name="ИП Тестов Тест Тестович",
    short_name="ИП Тестов Т.Т.",
    inn="123456789012",
    ogrnip="123456789012345",
    phone="+7 900 000-00-00",
    email="kp@example.ru",
    website="example.ru",
)


def content(**overrides) -> ProposalContent:
    d = ProposalDefaults()
    data = dict(
        customer_name="ООО «Клиника»",
        object_name="Медицинский центр",
        address="г. Москва, ул. Тестовая, 1",
        area_m2=D("812.5"),
        intro="Подготовили предложение.",
        sections=[
            ProposalSection(code="OBS", name="Обследование", description="Обмеры", stage="pre", price=D("100000")),
            ProposalSection(
                code="AR", name="Архитектурные решения (АР)", description="Планы", stage="rd", price=D("400000")
            ),
            ProposalSection(code="EOM", name="ЭОМ", stage="rd", price=D("300000")),
        ],
        total=D("800000"),
        duration=d.duration_text,
        payment_options=d.payment_options,
        vat_note=d.vat_note,
        quality=d.quality_text,
        advantages=d.advantages,
        signature="С уважением,\nкоманда",
        valid_until=date(2026, 11, 2),
        requisites=REQ,
    )
    data.update(overrides)
    return ProposalContent(**data)


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_money_format():
    assert money(D("1500000")) == "1 500 000 ₽"


def test_render_contains_customer_data_and_totals():
    text = docx_text(render_docx(content()))
    assert "ООО «Клиника»" in text
    assert "800 000 ₽" in text
    assert "Архитектурные решения (АР)" in text
    assert "Предпроектная стадия" in text and "Рабочая документация" in text
    assert "123456789012" in text
    assert "2 ноября 2026 г." in text
    assert "{{" not in text and "{%" not in text


def test_total_must_match_sections():
    with pytest.raises(ValueError):
        render_docx(content(total=D("1")))


def test_content_rejects_internal_fields():
    with pytest.raises(ValueError):
        content(margin=D("0.3"))


def test_user_text_is_escaped():
    text = docx_text(render_docx(content(customer_name="ООО <b>&</b>")))
    assert "ООО <b>&</b>" in text


WRITER_INSTALLED = shutil.which("soffice") is not None and any(
    Path(p).exists() for p in ("/usr/lib/libreoffice/program/libswlo.so", "/opt/libreoffice/program/libswlo.so")
)


@pytest.mark.skipif(not WRITER_INSTALLED, reason="LibreOffice Writer not installed")
def test_pdf_conversion():
    pdf = docx_to_pdf(render_docx(content()))
    assert pdf.startswith(b"%PDF")
