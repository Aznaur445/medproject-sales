"""Text extraction from tender documents: PDF, DOCX, XLSX, images (OCR via tesseract when installed).

Every chunk keeps its location (file, page/sheet) so extracted facts can point to where they were found.
"""

import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from app.core.logging import get_logger

log = get_logger(__name__)
MAX_PAGES = 200
PARAGRAPHS_PER_PAGE = 40  # DOCX has no pages: group paragraphs into pseudo-pages


@dataclass
class Chunk:
    file: str
    location: str  # «стр. 3», «лист "Смета"», «фрагмент 2»
    text: str


def _ocr_image(data: bytes) -> str:
    if shutil.which("tesseract") is None:
        return ""
    with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
        tmp.write(data)
        tmp.flush()
        result = subprocess.run(
            ["tesseract", tmp.name, "stdout", "-l", "rus+eng"], capture_output=True, timeout=180, check=False
        )
    return result.stdout.decode("utf-8", "ignore")


def _ocr_pdf(data: bytes, name: str) -> list[Chunk]:
    if shutil.which("pdftoppm") is None or shutil.which("tesseract") is None:
        return []
    chunks: list[Chunk] = []
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.pdf"
        src.write_bytes(data)
        subprocess.run(
            ["pdftoppm", "-r", "200", "-png", "-l", "30", str(src), str(Path(tmp) / "p")],
            capture_output=True,
            timeout=600,
            check=False,
        )
        for i, image in enumerate(sorted(Path(tmp).glob("p-*.png")), start=1):
            text = _ocr_image(image.read_bytes())
            if text.strip():
                chunks.append(Chunk(name, f"стр. {i} (скан)", text))
    return chunks


def extract_pdf(data: bytes, name: str) -> list[Chunk]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        chunks = [
            Chunk(name, f"стр. {i}", page.extract_text() or "")
            for i, page in enumerate(reader.pages[:MAX_PAGES], start=1)
        ]
    except Exception:  # noqa: BLE001 - broken PDFs happen; try OCR
        log.warning("pdf_parse_failed", file=name)
        chunks = []
    text_chunks = [c for c in chunks if len(c.text.strip()) > 20]
    if not text_chunks:  # a scan without a text layer
        return _ocr_pdf(data, name)
    return text_chunks


def extract_docx(data: bytes, name: str) -> list[Chunk]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_index, table in enumerate(doc.tables, start=1):
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(f"[таблица {t_index}] " + " | ".join(dict.fromkeys(cells)))
    return [
        Chunk(name, f"фрагмент {i // PARAGRAPHS_PER_PAGE + 1}", "\n".join(lines[i : i + PARAGRAPHS_PER_PAGE]))
        for i in range(0, len(lines), PARAGRAPHS_PER_PAGE)
    ]


def extract_xlsx(data: bytes, name: str) -> list[Chunk]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True, max_row=2000):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            chunks.append(Chunk(name, f'лист "{sheet.title}"', "\n".join(rows)))
    return chunks


def extract(data: bytes, name: str) -> list[Chunk]:
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            return extract_pdf(data, name)
        if suffix == ".docx":
            return extract_docx(data, name)
        if suffix in (".xlsx", ".xlsm"):
            return extract_xlsx(data, name)
        if suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            text = _ocr_image(data)
            return [Chunk(name, "изображение", text)] if text.strip() else []
        if suffix in (".txt", ".csv"):
            return [Chunk(name, "текст", data.decode("utf-8", "ignore"))]
    except Exception:  # noqa: BLE001
        log.exception("document_extract_failed", file=name)
    return []


def as_text(chunks: list[Chunk]) -> str:
    return "\n\n".join(f"[{c.file}, {c.location}]\n{c.text.strip()}" for c in chunks if c.text.strip())
